import csv
import hashlib
import io
import json
import re
import warnings
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError

from app.core.api import fail

MAX_ROWS = 10000
MAX_COLUMNS = 80
AD_PARSER_VERSION = 'amazon-ads-v2'
MONEY_FIELDS = ['item-price', 'item-tax', 'shipping-price', 'shipping-tax', 'gift-wrap-price', 'gift-wrap-tax',
                'item-promotion-discount', 'ship-promotion-discount']
SALES_REQUIRED = ['amazon-order-id', 'purchase-date', 'last-updated-date', 'order-status', 'fulfillment-channel',
                  'sales-channel', 'product-name', 'sku', 'asin', 'item-status', 'quantity', 'currency', 'item-price']
AD_COUNTS = {'展示量': 'impressions', '点击量': 'clicks', '7天总订单数(#)': 'orders', '7天总销售量(#)': 'units',
             '7天内广告SKU销售量(#)': 'advertised_units', '7天内其他SKU销售量(#)': 'other_units'}
AD_MONEY = {'花费': 'spend', '7天总销售额': 'attributed_sales', '7天内广告SKU销售额': 'advertised_sales', '7天内其他SKU销售额': 'other_sales'}
AD_TEXT = {'广告组合名称': 'portfolio', '货币': 'currency', '广告活动名称': 'campaign', '广告组名称': 'ad_group',
           '零售商': 'retailer', '国家/地区': 'country', '广告SKU': 'sku', '广告ASIN': 'asin'}
AD_REQUIRED = [*AD_TEXT, *AD_COUNTS, *AD_MONEY]


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()).hexdigest()


def clean(value, field, maximum=500, required=True):
    value = '' if value is None else str(value).strip()
    if (required and not value) or len(value) > maximum or any(ord(ch) < 32 for ch in value):
        raise ValueError(f'{field} 为空、过长或包含控制字符')
    return value


def amount(value, field, optional=False):
    if value is None or str(value).strip() == '':
        if optional:
            return None
        raise ValueError(f'{field} 不能为空')
    try:
        result = Decimal(str(value).replace(',', '').strip())
        if not result.is_finite() or result < 0 or result > Decimal('999999999999.9999'):
            raise ValueError(f'{field} 必须是范围内的非负金额')
        return format(result.quantize(Decimal('.0001'), rounding=ROUND_HALF_UP), 'f')
    except InvalidOperation as error:
        raise ValueError(f'{field} 金额格式无效') from error


def integer(value, field):
    try:
        number = Decimal(str(value).replace(',', '').strip())
        if not number.is_finite() or number < 0 or number > 1000000000 or number != number.to_integral_value():
            raise ValueError(f'{field} 必须为 0 到 10 亿之间的整数')
        return int(number)
    except InvalidOperation as error:
        raise ValueError(f'{field} 不是整数') from error


def timestamp(value, field):
    try:
        result = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        if result.tzinfo is None:
            raise ValueError()
        return result.astimezone(timezone.utc).isoformat()
    except (ValueError, TypeError, OverflowError) as error:
        raise ValueError(f'{field} 必须是带时区的时间') from error


def calendar_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    try:
        value = str(value).strip()
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}(?:[ T]00:00:00)?', value):
            raise ValueError()
        return date.fromisoformat(value[:10]).isoformat()
    except ValueError as error:
        raise ValueError('日期须为 YYYY-MM-DD 或 Excel 日期') from error


def sales_row(raw):
    data = {key.replace('-', '_'): clean(raw.get(key), key, 120 if key != 'product-name' else 2000)
            for key in ['amazon-order-id', 'sales-channel', 'sku', 'asin', 'order-status', 'item-status', 'fulfillment-channel', 'product-name']}
    data['purchase_date'] = timestamp(raw['purchase-date'], 'purchase-date')
    data['last_updated_date'] = timestamp(raw['last-updated-date'], 'last-updated-date')
    if data['last_updated_date'] < data['purchase_date']:
        raise ValueError('订单更新时间不能早于下单时间')
    data['quantity'] = integer(raw['quantity'], 'quantity')
    data['currency'] = clean(raw.get('currency'), 'currency', 3, required=False).upper() or None
    if data['currency'] and not re.fullmatch('[A-Z]{3}', data['currency']):
        raise ValueError('currency 须为三位币种代码')
    data.update({key.replace('-', '_'): amount(raw.get(key), key, optional=True) for key in MONEY_FIELDS})
    if any(data[key.replace('-', '_')] is not None for key in MONEY_FIELDS) and not data['currency']:
        raise ValueError('有金额的订单必须提供币种')
    if data['item_price'] is None:
        data['net_amount'] = None
    else:
        data['net_amount'] = format(sum((Decimal(data[key] or '0') * sign for key, sign in [
            ('item_price', 1), ('shipping_price', 1), ('gift_wrap_price', 1), ('item_promotion_discount', -1), ('ship_promotion_discount', -1)]), Decimal(0)), '.4f')
    data['merchant_order_id'] = clean(raw.get('merchant-order-id'), 'merchant-order-id', 120, required=False)
    key = fingerprint([data[field] for field in ['sales_channel', 'amazon_order_id', 'sku']])
    return key, key, data


def ad_keys(data):
    if data.get('report_date'):
        identity = fingerprint(['daily', data['sku'], data['campaign'], data['ad_group']])
        return fingerprint([identity, data['report_date']]), identity
    identity = fingerprint([data[field] for field in ['retailer', 'country', 'currency', 'campaign', 'ad_group', 'sku', 'asin']])
    return fingerprint([identity, data['start_date'], data['end_date']]), identity


def ad_row(raw):
    data = {field: clean(raw.get(label), label, 500 if field in {'portfolio', 'campaign', 'ad_group'} else 120,
                        required=field != 'portfolio') for label, field in AD_TEXT.items()}
    data['currency'] = data['currency'].upper()
    if not re.fullmatch('[A-Z]{3}', data['currency']):
        raise ValueError('货币须为三位币种代码')
    if '日期' in raw:
        day = calendar_date(raw['日期'])
        if any(raw.get(field) not in (None, '') and calendar_date(raw[field]) != day for field in ['开始日期', '结束日期']):
            raise ValueError('日期与开始/结束日期不一致，请核对日报')
        data['start_date'] = data['end_date'] = day
    else:
        data['start_date'], data['end_date'] = calendar_date(raw['开始日期']), calendar_date(raw['结束日期'])
    if data['start_date'] > data['end_date']:
        raise ValueError('开始日期不能晚于结束日期')
    if data['start_date'] == data['end_date']:
        data['report_date'] = data['start_date']
    data.update({field: integer(raw.get(label), label) for label, field in AD_COUNTS.items()})
    data.update({field: amount(raw.get(label), label) for label, field in AD_MONEY.items()})
    data['attribution_days'] = 7
    key, identity = ad_keys(data)
    return key, identity, data


def text_rows(content):
    try:
        encoding = 'utf-16' if content.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
        reader = csv.reader(io.StringIO(content.decode(encoding), newline=''), delimiter='\t', strict=True)
        yield from reader
    except (UnicodeError, csv.Error) as error:
        raise ValueError('销售文件须为 UTF-8 / UTF-16 编码的制表符 TXT / TSV，且列与引号有效') from error


def excel_rows(content):
    workbook = None
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(entry.file_size for entry in entries) > 50 * 1024 * 1024:
                raise ValueError('XLSX 解压后超过 50 MB 或文件结构过大')
            if any('vbaproject' in entry.filename.lower() for entry in entries):
                raise ValueError('请上传不含宏的原始 XLSX 报告')
        with warnings.catch_warnings():
            warnings.filterwarnings('ignore', message='Workbook contains no default style')
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if len(workbook.worksheets) != 1:
            raise ValueError('请使用仅包含一张报告数据表的 XLSX')
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()  # The supplied Amazon export incorrectly declares only A1.
        for cells in sheet.iter_rows(max_col=MAX_COLUMNS + 1):
            if any(cell.data_type == 'f' for cell in cells):
                yield ValueError('包含公式，请上传亚马逊直接导出的数值报告')
            else:
                row = [cell.value for cell in cells]
                while row and row[-1] is None:
                    row.pop()
                yield row
    except (zipfile.BadZipFile, KeyError, IndexError, InvalidFileException, ParseError, DefusedXmlException, EOFError, OSError) as error:
        raise ValueError('XLSX 文件损坏或格式不支持，请重新导出') from error
    finally:
        if workbook:
            workbook.close()


def parse_report(kind, content):
    if kind not in {'sales', 'ads'}:
        fail(422, 'invalid_report_type', '请选择销售或广告报告')
    rows, errors, duplicates, source_total, seen = [], [], 0, 0, {}
    generator = text_rows(content) if kind == 'sales' else excel_rows(content)
    try:
        headers = next(generator, None)
        if not isinstance(headers, list) or not headers:
            raise ValueError('文件为空或表头无效')
        headers = [str(value or '').strip() for value in headers]
        required = SALES_REQUIRED if kind == 'sales' else AD_REQUIRED
        if len(headers) > MAX_COLUMNS or len(headers) != len(set(headers)) or not set(required).issubset(headers):
            raise ValueError('表头不符合所选报告格式，缺少字段或包含重复列')
        if kind == 'ads' and '日期' not in headers and not {'开始日期', '结束日期'} <= set(headers):
            raise ValueError('广告报告须包含日期列，或开始日期与结束日期列')
        for index, values in enumerate(generator, 2):
            if index > MAX_ROWS + 1:
                raise ValueError(f'每次最多导入 {MAX_ROWS} 行，请分段导出')
            if isinstance(values, list) and not any(value is not None and str(value).strip() for value in values):
                continue
            source_total += 1
            try:
                if isinstance(values, ValueError):
                    raise values
                if kind == 'ads':
                    values += [None] * (len(headers) - len(values))
                if len(values) != len(headers):
                    raise ValueError('列数与表头不一致')
                raw = dict(zip(headers, values))
                key, identity, data = (sales_row if kind == 'sales' else ad_row)(raw)
                value_hash = fingerprint(data)
                if key in seen:
                    if kind == 'ads' and data.get('report_date'):
                        rows[seen[key][1]] = {'row': index, 'key': key, 'identity': identity, 'hash': value_hash, 'data': data}
                        seen[key] = (value_hash, seen[key][1])
                    elif seen[key][0] != value_hash:
                        raise ValueError('文件内同一业务键内容冲突，无法区分重复导出与同 SKU 拆行，请核对来源')
                    duplicates += 1
                    continue
                seen[key] = (value_hash, len(rows))
                rows.append({'row': index, 'key': key, 'identity': identity, 'hash': value_hash, 'data': data})
            except (ValueError, TypeError, OverflowError) as error:
                errors.append({'row': index, 'message': str(error)})
        if not source_total:
            raise ValueError('文件没有数据行')
    except ValueError as error:
        fail(422, 'invalid_report', str(error))
    finally:
        generator.close()
    return {'rows': rows, 'errors': errors, 'duplicate_count': duplicates, 'source_total': source_total}
